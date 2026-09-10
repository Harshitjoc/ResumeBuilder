from fastapi import APIRouter, HTTPException, UploadFile

router = APIRouter(prefix="/api/upload", tags=["upload"])

MAX_SIZE = 10 * 1024 * 1024  # 10 MB


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(__import__("io").BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


@router.post("/extract")
async def extract_resume_text(file: UploadFile):
    data = await file.read()
    if len(data) > MAX_SIZE:
        raise HTTPException(status_code=413, detail="File too large (max 10 MB)")

    ext = (file.filename or "").lower().rsplit(".", 1)[-1] if "." in (file.filename or "") else ""

    try:
        if ext == "pdf":
            text = _extract_pdf(data)
        elif ext == "docx":
            text = _extract_docx(data)
        elif ext == "txt":
            text = data.decode("utf-8", errors="replace").strip()
        else:
            raise HTTPException(status_code=415, detail="Unsupported file type (use .pdf, .docx, or .txt)")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not extract text: {e}")

    if not text:
        raise HTTPException(status_code=422, detail="No readable text found in the file (scan/OCR images not supported)")

    return {"filename": file.filename, "text": text}