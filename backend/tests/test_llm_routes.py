"""Real end-to-end smoke test for the LLM routes.

Runs the FastAPI app in-process via TestClient (no server process), hitting
Ollama on localhost with the small qwen3:0.6b model. Skips gracefully if
Ollama is unreachable.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

client = TestClient(app)

SAMPLE_RESUME = {
    "fullName": "Ada Lovelace",
    "contact": {"email": "ada@example.com", "phone": "+1 555 0100"},
    "summary": "Analytical engineer focused on computing systems.",
    "skills": ["Python", "Algorithms", "Analytics"],
    "workExperience": [
        {
            "company": "Analytical Engines",
            "title": "Engineer",
            "startDate": "2019-01",
            "endDate": "2022-06",
            "bullets": ["Built analytical models.", "Wrote algorithms in Python."],
        }
    ],
    "education": [
        {"school": "University of London", "degree": "BSc", "startDate": "2015-09", "endDate": "2018-06"}
    ],
    "projects": [],
    "certifications": [],
}

SAMPLE_JOB = """
Job Title: Data Analyst
Company: Acme Analytics
We need a data analyst with strong Python and SQL skills. You will analyze
datasets, build dashboards, and communicate findings to stakeholders.
2+ years experience required. Bonus for machine learning exposure.
"""

OLLAMA_KEYS = {"provider": "ollama", "apiKey": "http://127.0.0.1:11434", "model": "qwen3:0.6b"}


def check_ollama_up() -> bool:
    import urllib.request

    try:
        urllib.request.urlopen("http://127.0.0.1:11434", timeout=2)
        return True
    except Exception:
        return False


def test_bad_provider_rejected():
    resp = client.post(
        "/api/llm/generate-summary",
        json={"resume": SAMPLE_RESUME, "apiKeys": {"provider": "bogus"}},
    )
    assert resp.status_code == 400, resp.text


def test_missing_key_fails_cleanly():
    resp = client.post(
        "/api/llm/generate-summary",
        json={"resume": SAMPLE_RESUME, "apiKeys": {"provider": "openai"}},
    )
    assert resp.status_code == 502, resp.text


def test_generate_summary_ollama():
    resp = client.post(
        "/api/llm/generate-summary",
        json={"resume": SAMPLE_RESUME, "apiKeys": OLLAMA_KEYS},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert isinstance(body.get("summary", ""), str) and body["summary"].strip(), body
    print("SUMMARY:", body["summary"][:200])


def test_parse_job_ollama():
    resp = client.post(
        "/api/llm/parse-job",
        json={"jobText": SAMPLE_JOB, "apiKeys": OLLAMA_KEYS},
    )
    assert resp.status_code == 200, resp.text
    parsed = resp.json()["parsed"]
    assert parsed.get("job_title") or parsed.get("required_skills"), parsed
    print("PARSED JOB:", {
        k: parsed.get(k) for k in ("job_title", "company", "required_skills", "key_responsibilities")
    })


def test_parse_resume_ollama():
    raw_resume = """
    John A. Sample
    john.sample@email.com | 555-010-1234
    Senior Python Developer with 6 years of experience building data pipelines.

    SKILLS: Python, SQL, FastAPI, Docker, AWS

    EXPERIENCE
    Data Corp (2020 - present) - Senior Python Developer
    - Built ETL pipelines handling 2M rows daily.
    - Migrated services to FastAPI.

    EDUCATION
    BSc Computer Science, State University, 2016

    PROJECTS
    Resume Builder - personal project
    """
    resp = client.post(
        "/api/llm/parse-resume",
        json={"resumeText": raw_resume, "apiKeys": OLLAMA_KEYS},
    )
    assert resp.status_code == 200, resp.text
    parsed = resp.json()["parsed"]
    assert parsed.get("contact", {}).get("fullName"), parsed
    print("PARSED RESUME:", {
        "fullName": parsed.get("contact", {}).get("fullName"),
        "skills": parsed.get("skills"),
        "exp_count": len(parsed.get("experience", [])),
    })


def test_analyze_resume_ollama():
    resp = client.post(
        "/api/llm/analyze-resume",
        json={"resume": SAMPLE_RESUME, "apiKeys": OLLAMA_KEYS},
    )
    assert resp.status_code == 200, resp.text
    analysis = resp.json()["analysis"]
    assert isinstance(analysis.get("overall_score", 0), (int, float)), analysis
    assert isinstance(analysis.get("suggestions"), list), analysis
    print("ANALYZED RESUME:", {
        "score": analysis.get("overall_score"),
        "strengths": len(analysis.get("strengths", [])),
        "suggestions": len(analysis.get("suggestions", [])),
    })


def test_extract_docx():
    import io

    import docx

    d = docx.Document()
    d.add_paragraph("Grace Hopper")
    d.add_paragraph("grace@navy.mil")
    d.add_paragraph("SKILLS: COBOL, Mathematics")
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/upload/extract",
        files={"file": ("resume.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    text = resp.json()["text"]
    assert "Grace Hopper" in text and "COBOL" in text, text
    print("EXTRACTED DOCX:", text.replace("\n", " / ")[:120])


def test_extract_unsupported_type():
    resp = client.post("/api/upload/extract", files={"file": ("cv.exe", b"MZ...", "application/octet-stream")})
    assert resp.status_code == 415, resp.text


if __name__ == "__main__":
    failures = 0
    tests = [
        test_bad_provider_rejected,
        test_missing_key_fails_cleanly,
        test_extract_docx,
        test_extract_unsupported_type,
    ]
    if check_ollama_up():
        tests += [
            test_generate_summary_ollama,
            test_parse_job_ollama,
            test_parse_resume_ollama,
            test_analyze_resume_ollama,
        ]
    else:
        print("SKIP: Ollama not reachable, running error/file tests only")

    for t in tests:
        print(f"== {t.__name__}")
        try:
            t()
            print("   PASS")
        except AssertionError as e:
            failures += 1
            print(f"   FAIL: {e}")
        except Exception as e:
            failures += 1
            print(f"   ERROR: {type(e).__name__}: {e}")

    print(f"\n{'ALL PASSED' if failures == 0 else f'{failures} FAILURE(S)'}")
    sys.exit(1 if failures else 0)